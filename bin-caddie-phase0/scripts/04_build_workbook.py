"""Step 4: assemble the Phase 0 working file.

Inputs : outputs/hosts_ranked.csv, outputs/council_da_targets.csv, outputs/drop_plan.csv,
         data/raw/accounts_batch_*.csv (published business contacts for the top 40 hosts)
Output : outputs/Phase0_Workbook.xlsx

Tabs: Dashboard · Accounts · Outreach Log · Suppression · Letterbox Drops ·
      Council DA Targets · Noosa Sign Log · Tier B Hosts · Read Me
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT, RAW = ROOT / "outputs", ROOT / "data" / "raw"

NOOSA = ("Noosa Heads", "Noosaville", "Sunshine Beach", "Tewantin", "Noosa Hinterland", "Peregian")
HEAD = PatternFill("solid", fgColor="16324F")
SUB = PatternFill("solid", fgColor="EAF6EF")
BOLD_W = Font(bold=True, color="FFFFFF")
OUTCOMES = ["no reply", "replied", "meeting booked", "pilot agreed", "not interested", "unsubscribe", "call back later"]


def angle_for(row) -> str:
    first_suburb = str(row.top_suburbs).split(" (")[0]
    if first_suburb.startswith(NOOSA):
        return "B - Noosa contact person"
    if row.units > row.houses:
        return "C - owner-services add-on / bin rooms"
    return "A - guest bins nobody owns"


def load_accounts(hosts: pd.DataFrame) -> pd.DataFrame:
    batches = sorted(RAW.glob("accounts_batch_*.csv"))
    researched = pd.concat([pd.read_csv(p, dtype=str) for p in batches], ignore_index=True) if batches else pd.DataFrame()
    a = hosts[hosts.active_listings_here >= 10].copy()
    a["angle"] = a.apply(angle_for, axis=1)
    if not researched.empty:
        researched.columns = [c.strip().lower() for c in researched.columns]
        a = a.merge(researched, on="host_name", how="left", suffixes=("", "_r"))
    for col in ["business_name", "website", "business_type", "office_suburb", "published_email",
                "email_source_url", "published_phone", "phone_source_url", "decision_maker",
                "no_marketing_notice", "portfolio_notes", "current_bin_or_cleaning_mentions",
                "pitch_angle", "confidence", "status"]:
        if col not in a:
            a[col] = ""
    a = a.fillna("")
    has_email = a.published_email.str.contains("@")
    has_phone = a.published_phone.str.contains(r"\d{6}", regex=True) | a.published_phone.str.replace(" ", "").str.contains(r"\d{8}")
    contactable = (
        (has_email | has_phone)
        & ~a.no_marketing_notice.str.lower().eq("yes")
        & ~a.confidence.str.lower().eq("low")
        & ~a.status.str.lower().str.contains("individual|unidentified|do not contact")
    )
    a["contactable"] = contactable.map({True: "yes", False: "no"})
    # Phone-only accounts start with a business-landline call, then a contact-form enquiry.
    # Researchers flag emails that are the wrong inbox (sister brand, guest-only) in `status`.
    email_ok = has_email & ~a.status.str.contains(r"phone only|call office|phone/form|phone/web form|partner form", case=False)
    a["first_channel"] = ""
    a.loc[contactable & email_ok, "first_channel"] = "email"
    a.loc[contactable & ~email_ok, "first_channel"] = "phone (business line) / web form"

    # Several Airbnb host names are brands of one company: pitch the company once.
    a["group"] = a.business_name.where(a.business_name != "", a.host_name)
    a.loc[a.business_name.str.contains("Alloggio", case=False), "group"] = "Alloggio Group"
    a.loc[a.business_name.str.contains("Utopia", case=False), "group"] = "Stay Utopia"
    a["group_listings_total"] = a.groupby("group").active_listings_here.transform("sum")

    # Phase 0 test group: 5 groups per angle, biggest portfolios first, one row per group.
    a["test_group"] = ""
    lead_rows = a[contactable].sort_values("active_listings_here", ascending=False).drop_duplicates("group")
    for ang in a.angle.unique():
        idx = lead_rows[lead_rows.angle == ang].sort_values("group_listings_total", ascending=False).index[:5]
        a.loc[idx, "test_group"] = "Phase 0 wave 1"
    cols = ["test_group", "group", "group_listings_total", "host_name", "business_name",
            "active_listings_here", "houses", "units",
            "top_suburbs", "angle", "pitch_angle", "first_channel", "published_email", "email_source_url", "published_phone",
            "phone_source_url", "decision_maker", "website", "business_type", "office_suburb",
            "no_marketing_notice", "contactable", "confidence", "status", "portfolio_notes",
            "current_bin_or_cleaning_mentions"]
    return a.sort_values(["test_group", "active_listings_here"], ascending=[False, False])[cols]


def header(ws, row: int, cols: list[str]) -> None:
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font = HEAD, BOLD_W
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def write_df(ws, df: pd.DataFrame, start_row: int = 1, widths: dict | None = None) -> None:
    header(ws, start_row, list(df.columns))
    for r, rec in enumerate(df.itertuples(index=False), start_row + 1):
        for c, v in enumerate(rec, 1):
            ws.cell(row=r, column=c, value=None if (isinstance(v, float) and pd.isna(v)) else v)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=3)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}"
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, min(40, max(12, len(col) + 2)))


def blank_log(ws, cols: list[str], rows: int = 300) -> None:
    header(ws, 1, cols)
    ws.freeze_panes = "A2"
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(c) + 4)


def main() -> None:
    hosts = pd.read_csv(OUT / "hosts_ranked.csv")
    das = pd.read_csv(OUT / "council_da_targets.csv")
    drops = pd.read_csv(OUT / "drop_plan.csv")
    accounts = load_accounts(hosts)

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"

    # ---- Accounts --------------------------------------------------------
    ws = wb.create_sheet("Accounts")
    write_df(ws, accounts, widths={"host_name": 28, "business_name": 28, "top_suburbs": 36, "angle": 30,
                                   "pitch_angle": 50, "published_email": 30, "email_source_url": 36,
                                   "portfolio_notes": 50, "test_group": 16})
    ws.conditional_formatting.add(f"A2:A{len(accounts) + 1}", CellIsRule(operator="equal",
                                  formula=['"Phase 0 wave 1"'], fill=SUB))

    # ---- Outreach log ----------------------------------------------------
    ws = wb.create_sheet("Outreach Log")
    blank_log(ws, ["date", "account (host_name)", "contact role", "angle", "channel", "touch #",
                   "outcome", "next step", "next date", "notes"])
    dv = DataValidation(type="list", formula1='"' + ",".join(OUTCOMES) + '"', allow_blank=True)
    ch = DataValidation(type="list", formula1='"email,phone,mail,in person,referral"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(ch)
    dv.add("G2:G1000")
    ch.add("E2:E1000")

    # ---- Suppression -----------------------------------------------------
    ws = wb.create_sheet("Suppression")
    blank_log(ws, ["date", "email / phone / business", "channel", "how they opted out", "logged by"])
    ws["G1"] = "Anyone here is NEVER contacted again, on any channel. Action unsubscribes within 5 working days (aim: same day)."
    ws["G1"].font = Font(bold=True, color="C0392B")

    # ---- Letterbox drops -------------------------------------------------
    ws = wb.create_sheet("Letterbox Drops")
    d = drops.copy()
    for c in ["date dropped", "letterboxes delivered", "print cost $", "other cost $",
              "QR visits (GA4)", "calls/forms quoting code", "leads", "customers won"]:
        d[c] = None
    write_df(ws, d, widths={"qr_url": 60, "centre_map_link": 40})
    n = len(d)
    cols = {c: get_column_letter(i + 1) for i, c in enumerate(d.columns)}
    cpl_col = get_column_letter(len(d.columns) + 1)
    cac_col = get_column_letter(len(d.columns) + 2)
    for col, title in ((cpl_col, "cost per lead $"), (cac_col, "cost per customer $")):
        cell = ws[f"{col}1"]
        cell.value, cell.fill, cell.font = title, HEAD, BOLD_W
        ws.column_dimensions[col].width = 18
    for r in range(2, n + 2):
        cost = f"(N({cols['print cost $']}{r})+N({cols['other cost $']}{r}))"
        ws[f"{cpl_col}{r}"] = f'=IF(N({cols["leads"]}{r})=0,"",{cost}/{cols["leads"]}{r})'
        ws[f"{cac_col}{r}"] = f'=IF(N({cols["customers won"]}{r})=0,"",{cost}/{cols["customers won"]}{r})'

    # ---- Council DA targets ----------------------------------------------
    ws = wb.create_sheet("Council DA Targets")
    da = das.copy()
    da["date_received"] = pd.to_datetime(da.date_received).dt.date
    da["letter sent"] = None
    da["response"] = None
    write_df(ws, da[["signal", "pitch", "date_received", "address", "use", "status", "decision",
                     "ram_id", "hex", "letter sent", "response"]],
             widths={"address": 44, "use": 60, "pitch": 24})

    # ---- Noosa sign log --------------------------------------------------
    ws = wb.create_sheet("Noosa Sign Log")
    blank_log(ws, ["date", "street", "house no.", "suburb", "approval no.", "contact type",
                   "agency name (business only)", "notes"])
    ct = DataValidation(type="list", formula1='"agency,individual"', allow_blank=True)
    ws.add_data_validation(ct)
    ct.add("F2:F2000")
    ws["J1"] = "Individuals: record 'individual' only - never their name or number (complaints line, not marketing)."
    ws["J1"].font = Font(bold=True, color="C0392B")

    # ---- Tier B ----------------------------------------------------------
    ws = wb.create_sheet("Tier B Hosts")
    tb = hosts[(hosts.active_listings_here >= 5) & (hosts.active_listings_here < 10)][
        ["host_name", "active_listings_here", "houses", "units", "top_suburbs", "looks_like_business", "outreach_route"]]
    write_df(ws, tb, widths={"host_name": 30, "top_suburbs": 40, "outreach_route": 50})

    # ---- Dashboard -------------------------------------------------------
    dash["A1"] = "Bin Caddie - Phase 0 scorecard"
    dash["A1"].font = Font(bold=True, size=16, color="16324F")
    dash["A2"] = "Pass = 3+ manager meetings AND 1+ pilot signed. Outcomes are counted from the Outreach Log."
    contact_col = get_column_letter(list(accounts.columns).index("contactable") + 1)
    rows = [
        ("Wave 1 accounts selected", "=COUNTIF(Accounts!A:A,\"Phase 0 wave 1\")", None),
        ("Accounts contactable (published business email/phone)", f"=COUNTIF(Accounts!{contact_col}:{contact_col},\"yes\")", None),
        ("Touches logged", "=COUNTA('Outreach Log'!A:A)-1", None),
        ("Replies", "=COUNTIF('Outreach Log'!G:G,\"replied\")+COUNTIF('Outreach Log'!G:G,\"meeting booked\")+COUNTIF('Outreach Log'!G:G,\"pilot agreed\")", None),
        ("Meetings booked", "=COUNTIF('Outreach Log'!G:G,\"meeting booked\")+COUNTIF('Outreach Log'!G:G,\"pilot agreed\")", 3),
        ("Pilots agreed", "=COUNTIF('Outreach Log'!G:G,\"pilot agreed\")", 1),
        ("Unsubscribes", "=COUNTA(Suppression!A:A)-1", None),
        ("Letterboxes delivered", "=SUM('Letterbox Drops'!L:L)", 2000),
        ("Letterbox leads", "=SUM('Letterbox Drops'!Q:Q)", None),
        ("Letterbox customers", "=SUM('Letterbox Drops'!R:R)", None),
        ("Noosa signs logged", "=COUNTA('Noosa Sign Log'!A:A)-1", None),
        ("Council 'To the Owner' letters sent", "=COUNTA('Council DA Targets'!J:J)-1", None),
    ]
    header(dash, 4, ["Metric", "Actual", "Target", "Status"])
    for i, (label, formula, target) in enumerate(rows, 5):
        dash.cell(row=i, column=1, value=label)
        dash.cell(row=i, column=2, value=formula)
        if target is not None:
            dash.cell(row=i, column=3, value=target)
            dash.cell(row=i, column=4, value=f'=IF(B{i}>=C{i},"PASS","in progress")')
    dash.column_dimensions["A"].width = 42
    for c in "BCD":
        dash.column_dimensions[c].width = 14
    dash.conditional_formatting.add("D5:D20", CellIsRule(operator="equal", formula=['"PASS"'], fill=SUB))

    # ---- Read me ---------------------------------------------------------
    ws = wb.create_sheet("Read Me")
    notes = [
        "Sources: Inside Airbnb Sunshine Coast 28 Sep 2026 (internal analysis only - do not share or republish);",
        "Sunshine Coast Council DA open data (CC BY 3.0 AU) - names stripped, letters go to 'The Owner';",
        "Account contacts: published business details only, with source URL (Spam Act conspicuous publication).",
        "",
        "Rules: no Airbnb messaging or scraping; no cold SMS; business landlines 9am-8pm weekdays, 9am-5pm Sat, never Sun;",
        "wash mobiles against the Do Not Call Register; every email identifies Bin Caddie and has an unsubscribe;",
        "never contact individual hosts by name; never mention bins being left out.",
        "",
        "Unverified: Noosa bin/kerbside condition wording - check current approval conditions before using in copy.",
        "Offers marked (confirm) in the templates need Bin Caddie's sign-off (trial size, portfolio pricing).",
    ]
    for i, t in enumerate(notes, 1):
        ws.cell(row=i, column=1, value=t)
    ws.column_dimensions["A"].width = 120

    wb.save(OUT / "Phase0_Workbook.xlsx")
    print(f"Accounts: {len(accounts)} (wave 1: {(accounts.test_group != '').sum()}, "
          f"contactable: {(accounts.contactable == 'yes').sum()}) -> outputs/Phase0_Workbook.xlsx")


if __name__ == "__main__":
    main()
